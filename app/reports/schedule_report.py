"""Export des emplois du temps en VRAIS fichiers PDF (reportlab) et Excel (openpyxl).

Le rendu reprend l'identité visuelle de l'application :
  - mêmes couleurs (constantes C et PALETTE), même police (Segoe UI si disponible),
  - fond de page gris-bleu (C["bg"]) avec une « carte » blanche arrondie, comme les pages de l'application,
  - bandeau sombre (C["sidebar"]) avec filet primaire, comme la barre latérale,
  - colonnes = jours (le jour courant est mis en évidence en couleur primaire), lignes = créneaux,
  - cases arrondies colorées par matière, cases vides en gris clair, pauses sur toute la largeur.

Interface inchangée : build_tables(...) puis write_file(path, tables, fmt) avec fmt = "pdf" | "xlsx".
Dans pages/schedule.py : `from app.services import schedule_export as export_service`.
"""
import datetime
import os
import re

from app.constants import C, DAYS, FONT, PALETTE, PLAN, SLOTS
from app.services.schedule_service import cell_data
from app.services.settings_service import get_settings
from app.utils.formatting import ordinal


# ----------------------------------------------------------------------------- données
def _safe(name):
    return re.sub(r"[^\w\-]+", "_", name).strip("_") or "emploi_du_temps"


def _colors(key):
    bg, fg = PALETTE[key % len(PALETTE)]
    return bg, fg


def _week(mode, ident, name):
    data = cell_data(mode, ident)
    if mode == "c":
        return dict(kind="week", mode="c", sheet=name, title=f"Emploi du temps — {name}", subtitle="", cells=data)
    n = len(data)
    return dict(kind="week", mode="t", sheet=name, title=name, subtitle=f"{n} heure{'s' if n > 1 else ''} par semaine",
                cells=data)


def _general(cls_map):
    classes, legend = [], {}
    for name, cid in cls_map.items():
        data = cell_data("c", cid)
        classes.append((name, data))
        for x in data.values():
            legend.setdefault(x["code"], (x["subj"], x["key"]))
    return dict(kind="general", sheet="Général", title="Emploi du temps général", subtitle="", classes=classes,
                legend=dict(sorted(legend.items())))


def build_tables(scope, cid, tid, cls_map, teachers):
    """-> (liste de tableaux, nom de fichier de base). Liste vide s'il n'y a rien à exporter."""
    rcls = {i: n for n, i in cls_map.items()}
    rtea = {i: n for n, i in teachers.items()}
    if scope == "classe":
        return [_week("c", cid, rcls.get(cid, ""))], "Emploi_du_temps_" + _safe(rcls.get(cid, "classe"))
    if scope == "enseignant":
        return [_week("t", tid, rtea.get(tid, ""))], "Emploi_du_temps_" + _safe(rtea.get(tid, "enseignant"))
    if scope == "classes":
        return [_week("c", i, n) for n, i in cls_map.items()], "Emplois_du_temps_classes"
    if scope == "enseignants":
        tabs = [_week("t", i, n) for n, i in teachers.items()]
        return [t for t in tabs if t["cells"]], "Emplois_du_temps_enseignants"
    return ([_general(cls_map)] if cls_map else []), "Emploi_du_temps_general"


def write_file(path, tables, fmt):
    if fmt == "pdf":
        return _write_pdf(path, tables)
    if fmt == "xlsx":
        return _write_xlsx(path, tables)
    raise ValueError(f"Format non pris en charge : {fmt}")


def _lines(mode, x):
    """Les deux lignes affichées dans une case, comme dans le calendrier de l'application."""
    return (x["subj"], x["teacher"]) if mode == "c" else (x["cls"], x["subj"])


# ============================================================================= PDF
def _register_fonts():
    """Enregistre Segoe UI (police de l'application) si présente ; sinon Helvetica."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    fonts = os.environ.get("WINDIR", r"C:\Windows") + r"\Fonts"
    reg, bold = os.path.join(fonts, "segoeui.ttf"), os.path.join(fonts, "segoeuib.ttf")
    try:
        if os.path.exists(reg) and os.path.exists(bold):
            pdfmetrics.registerFont(TTFont("SegoeUI", reg))
            pdfmetrics.registerFont(TTFont("SegoeUI-Bold", bold))
            return "SegoeUI", "SegoeUI-Bold"
    except Exception:
        pass
    return "Helvetica", "Helvetica-Bold"


def _write_pdf(path, tables):
    from reportlab.lib.colors import HexColor
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.pdfgen import canvas

    F, FB = _register_fonts()
    PW, PH = landscape(A4)
    M = 28                                   # marge de la page
    PAD = 16                                 # marge intérieure de la carte
    X0 = M + PAD                             # début du contenu
    W = PW - 2 * X0                          # largeur du contenu
    BANNER_H = 40
    S = get_settings()
    today = datetime.date.today().strftime("%d/%m/%Y")
    today_d = datetime.date.today().weekday()
    col = lambda h: HexColor(h)
    TEXT, MUTED, BORDER, SOFT = col(C["text"]), col(C["muted"]), col(C["border"]), col(C["soft"])
    PRIMARY, SIDEBAR = col(C["primary"]), col(C["sidebar"])
    PILL, BAR = col("#F1F5F9"), col("#EEF2F7")

    card_top = PH - BANNER_H - 14
    card_bottom = M
    top_y = card_top - 56                    # début du tableau (sous le titre)
    avail = top_y - (card_bottom + 14)

    def fit(text, font, size, maxw):
        text = str(text or "")
        if stringWidth(text, font, size) <= maxw:
            return text
        while text and stringWidth(text + "…", font, size) > maxw:
            text = text[:-1]
        return text + "…"

    def wrap2(text, font, size, maxw):
        """Texte sur 1 ou 2 lignes (la 2e est tronquée avec « … » si besoin)."""
        text = str(text or "")
        if stringWidth(text, font, size) <= maxw:
            return [text]
        words, first = text.split(), ""
        for i, w in enumerate(words):
            nxt = (first + " " + w).strip()
            if stringWidth(nxt, font, size) > maxw and first:
                return [first, fit(" ".join(words[i:]), font, size, maxw)]
            first = nxt
        return [fit(text, font, size, maxw)]

    # ---- pagination (pour « Page n / N »)
    pages = []
    ROW_H, HEAD_H = 16, 28
    for t in tables:
        if t["kind"] == "week":
            pages.append((t, 0, None, False))
            continue
        block = len(DAYS) * ROW_H + 5
        per_page = max(1, int((avail - HEAD_H) // block))
        names = list(range(len(t["classes"])))
        chunks = [names[i:i + per_page] for i in range(0, len(names), per_page)] or [[]]
        leg_h = 22 + 15 * (-(-len(t["legend"]) // 3))
        for k, ch in enumerate(chunks):
            last = k == len(chunks) - 1
            room = avail - HEAD_H - len(ch) * block
            pages.append((t, k, ch, last and room >= leg_h))
        if not pages[-1][3]:
            pages.append((t, len(chunks), [], True))
    total = len(pages)

    c = canvas.Canvas(path, pagesize=(PW, PH), pageCompression=1)
    c.setTitle(tables[0]["title"] if len(tables) == 1 else "Emplois du temps")
    c.setAuthor("EduManager")

    def cell(x, y, w, h, bg, edge, radius=7):
        c.setFillColor(col(bg) if isinstance(bg, str) else bg)
        c.setStrokeColor(col(edge) if isinstance(edge, str) else edge)
        c.setLineWidth(0.7)
        c.roundRect(x, y, w, h, radius, fill=1, stroke=1)

    def page_frame(t, no):
        # fond de page (comme C["bg"])
        c.setFillColor(col(C["bg"]))
        c.rect(0, 0, PW, PH, fill=1, stroke=0)
        # bandeau sombre (comme la barre latérale) + filet primaire
        c.setFillColor(SIDEBAR)
        c.rect(0, PH - BANNER_H, PW, BANNER_H, fill=1, stroke=0)
        c.setFillColor(PRIMARY)
        c.rect(0, PH - BANNER_H - 2.5, PW, 2.5, fill=1, stroke=0)
        c.setFillColor(col("#FFFFFF"))
        c.setFont(FB, 11)
        c.drawString(M, PH - BANNER_H / 2 - 4, str(S["ecole_nom"]).upper())
        c.setFillColor(col(C["sidebar_text"]))
        c.setFont(F, 9)
        c.drawRightString(PW - M, PH - BANNER_H / 2 - 3, f"Année scolaire {S['annee']}")
        # carte blanche arrondie
        cell(M, card_bottom, PW - 2 * M, card_top - card_bottom, C["card"], C["border"], 12)
        # titre
        c.setFillColor(TEXT)
        c.setFont(FB, 17)
        c.drawString(X0, card_top - 28, fit(t["title"], FB, 17, W))
        if t.get("subtitle"):
            c.setFillColor(MUTED)
            c.setFont(F, 10)
            c.drawString(X0, card_top - 43, t["subtitle"])
        # pied de page
        c.setFillColor(MUTED)
        c.setFont(F, 7.5)
        c.drawString(M, M - 14, f"Édité le {today}")
        c.drawRightString(PW - M, M - 14, f"Page {no} / {total}")

    # ---- vue hebdomadaire (classe ou enseignant)
    def draw_week(t):
        time_w, gap = 64, 3
        col_w = (W - time_w) / len(DAYS)
        n_breaks = sum(1 for p in PLAN if p[0] != "slot")
        head_h, break_h = 24, 17
        slot_h = min(66, (avail - head_h - n_breaks * break_h - gap * (len(PLAN) + 1)) / max(1, len(SLOTS)))
        y = top_y
        for d, name in enumerate(DAYS):                              # en-têtes de jours
            x = X0 + time_w + d * col_w + gap / 2
            hl = d == today_d
            cell(x, y - head_h, col_w - gap, head_h - 2, PRIMARY if hl else PILL, PRIMARY if hl else PILL, 8)
            c.setFillColor(col("#FFFFFF") if hl else TEXT)
            c.setFont(FB, 10)
            c.drawCentredString(x + (col_w - gap) / 2, y - head_h / 2 - 4, name)
        y -= head_h + gap
        si = 0
        for kind, lab, s, e in PLAN:
            if kind == "slot":
                c.setFillColor(MUTED)
                c.setFont(FB, 8.5)
                c.drawCentredString(X0 + time_w / 2, y - slot_h / 2 + 2, ordinal(si))
                c.setFont(F, 7.5)
                c.drawCentredString(X0 + time_w / 2, y - slot_h / 2 - 8, f"{s}-{e}")
                for d in range(len(DAYS)):
                    x = X0 + time_w + d * col_w + gap / 2
                    w, h = col_w - gap, slot_h - gap
                    xc = t["cells"].get((d, si))
                    if xc:
                        bg, fg = _colors(xc["key"])
                        cell(x, y - slot_h + gap / 2, w, h, bg, bg, 9)
                        l1, l2 = _lines(t["mode"], xc)
                        lines = [(ln, FB, 9.5) for ln in wrap2(l1, FB, 9.5, w - 16)]
                        if l2:
                            lines.append((fit(l2, F, 8.5, w - 16), F, 8.5))
                        block_h = 11.5 * len(lines)
                        yy = y - slot_h / 2 + block_h / 2 - 9 + gap / 2
                        c.setFillColor(col(fg))
                        for text, font, size in lines:
                            c.setFont(font, size)
                            c.drawString(x + 8, yy, text)
                            yy -= 11.5
                    else:
                        cell(x, y - slot_h + gap / 2, w, h, SOFT, BORDER, 9)
                y -= slot_h + gap
                si += 1
            else:
                cell(X0 + gap / 2, y - break_h, W - gap, break_h - 2, BAR, BAR, 6)
                c.setFillColor(MUTED)
                c.setFont(FB, 8)
                c.drawCentredString(X0 + W / 2, y - break_h / 2 - 3, f"{lab}  ·  {s} – {e}")
                y -= break_h + gap

    # ---- vue générale (toutes les classes)
    def draw_general(t, chunk, with_legend):
        cls_w, day_w, gap = 62, 52, 1.5
        sl_w = (W - cls_w - day_w) / len(SLOTS)
        y = top_y
        if chunk is not None:
            for i in range(len(SLOTS) + 2):                          # en-tête primaire
                if i == 0:
                    x, w, txt = X0, cls_w, "CLASSE"
                elif i == 1:
                    x, w, txt = X0 + cls_w, day_w, "JOURS"
                else:
                    x, w, txt = X0 + cls_w + day_w + (i - 2) * sl_w, sl_w, None
                cell(x + gap, y - HEAD_H, w - 2 * gap, HEAD_H - 2, PRIMARY, PRIMARY, 5)
                c.setFillColor(col("#FFFFFF"))
                c.setFont(FB, 8.5)
                if txt:
                    c.drawCentredString(x + w / 2, y - HEAD_H / 2 - 3, txt)
                else:
                    c.drawCentredString(x + w / 2, y - 11, ordinal(i - 2))
                    c.setFont(F, 6.5)
                    s, e = SLOTS[i - 2]
                    c.drawCentredString(x + w / 2, y - 21, f"{s}-{e}")
            y -= HEAD_H + 3
            for ci in chunk:
                name, data = t["classes"][ci]
                cell(X0 + gap, y - len(DAYS) * ROW_H + gap, cls_w - 2 * gap, len(DAYS) * ROW_H - 2 * gap, BAR, BAR, 6)
                c.setFillColor(TEXT)
                c.setFont(FB, 9)
                c.drawCentredString(X0 + cls_w / 2, y - len(DAYS) * ROW_H / 2 - 3, fit(name, FB, 9, cls_w - 8))
                for d, day in enumerate(DAYS):
                    ry = y - (d + 1) * ROW_H
                    cell(X0 + cls_w + gap, ry + gap, day_w - 2 * gap, ROW_H - 2 * gap, PILL, PILL, 4)
                    c.setFillColor(MUTED)
                    c.setFont(FB, 7.5)
                    c.drawCentredString(X0 + cls_w + day_w / 2, ry + ROW_H / 2 - 2.5, day)
                    for s in range(len(SLOTS)):
                        x = X0 + cls_w + day_w + s * sl_w
                        xc = data.get((d, s))
                        if xc:
                            bg, fg = _colors(xc["key"])
                            cell(x + gap, ry + gap, sl_w - 2 * gap, ROW_H - 2 * gap, bg, bg, 4)
                            c.setFillColor(col(fg))
                            c.setFont(FB, 7.5)
                            c.drawCentredString(x + sl_w / 2, ry + ROW_H / 2 - 2.5, fit(xc["code"], FB, 7.5, sl_w - 6))
                        else:
                            cell(x + gap, ry + gap, sl_w - 2 * gap, ROW_H - 2 * gap, SOFT, BORDER, 4)
                y -= len(DAYS) * ROW_H + 5
        if with_legend and t["legend"]:
            c.setFillColor(TEXT)
            c.setFont(FB, 10)
            c.drawString(X0, y - 12, "Légende")
            y -= 22
            col_w = W / 3
            for k, (code, (subj, key)) in enumerate(t["legend"].items()):
                x, yy = X0 + (k % 3) * col_w, y - (k // 3) * 15
                bg, fg = _colors(key)
                cell(x, yy - 11, 38, 12, bg, bg, 3)
                c.setFillColor(col(fg))
                c.setFont(FB, 7.5)
                c.drawCentredString(x + 19, yy - 8, fit(code, FB, 7.5, 34))
                c.setFillColor(TEXT)
                c.setFont(F, 8.5)
                c.drawString(x + 44, yy - 8.5, fit(subj, F, 8.5, col_w - 52))

    for no, (t, k, chunk, legend) in enumerate(pages, start=1):
        page_frame(t, no)
        if t["kind"] == "week":
            draw_week(t)
        else:
            draw_general(t, chunk if chunk else None, legend)
        c.showPage()
    c.save()
    return path


# ============================================================================= Excel
def _write_xlsx(path, tables):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.properties import PageSetupProperties
    try:
        from openpyxl.cell.rich_text import CellRichText, TextBlock
        from openpyxl.cell.text import InlineFont
    except ImportError:                                  # openpyxl < 3.1 : texte simple (matière en gras)
        CellRichText = None

    S = get_settings()
    FN = FONT                                            # « Segoe UI », comme l'application
    today_d = datetime.date.today().weekday()
    hx = lambda h: h.lstrip("#").upper()
    fill = lambda h: PatternFill("solid", fgColor=hx(h))
    gap = Side(style="medium", color="FFFFFF")           # liseré blanc = espace entre les cases, comme à l'écran
    gapb = Border(left=gap, right=gap, top=gap, bottom=gap)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    FIRST = 6                                            # première ligne du tableau

    def sheet(name):
        base = re.sub(r"[\[\]\*\?/\\:]", " ", name or "Feuille").strip()[:28] or "Feuille"
        n, k = base, 2
        while n.lower() in used:
            n, k = f"{base[:26]} {k}", k + 1
        used.add(n.lower())
        ws = wb.create_sheet(n)
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.page_margins.top = ws.page_margins.bottom = 0.5
        ws.sheet_properties.tabColor = hx(C["primary"])
        return ws

    def title_rows(ws, t, ncols):
        # ligne 1 : bandeau sombre (comme la barre latérale)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
        b = ws.cell(1, 1, f"{str(S['ecole_nom']).upper()}   ·   Année scolaire {S['annee']}")
        b.font = Font(name=FN, size=10, bold=True, color="FFFFFF")
        b.alignment = Alignment(vertical="center", indent=1)
        for cc in range(1, ncols + 1):
            ws.cell(1, cc).fill = fill(C["sidebar"])
        ws.row_dimensions[1].height = 26
        # ligne 2 : filet primaire
        for cc in range(1, ncols + 1):
            ws.cell(2, cc).fill = fill(C["primary"])
        ws.row_dimensions[2].height = 3
        # ligne 3 : titre ; ligne 4 : sous-titre
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=ncols)
        ws.cell(3, 1, t["title"]).font = Font(name=FN, size=16, bold=True, color=hx(C["text"]))
        ws.cell(3, 1).alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[3].height = 32
        ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=ncols)
        ws.cell(4, 1, t.get("subtitle") or "").font = Font(name=FN, size=10, color=hx(C["muted"]))
        ws.cell(4, 1).alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[4].height = 16
        ws.row_dimensions[5].height = 8

    # ---- vue hebdomadaire
    def week(t):
        ws = sheet(t["sheet"])
        ncols = len(DAYS) + 1
        title_rows(ws, t, ncols)
        ws.column_dimensions["A"].width = 15
        for d in range(len(DAYS)):
            ws.column_dimensions[get_column_letter(d + 2)].width = 25
        r = FIRST
        ws.row_dimensions[r].height = 26
        for d, name in enumerate(DAYS):
            hl = d == today_d
            cl = ws.cell(r, d + 2, name)
            cl.fill = fill(C["primary"] if hl else "#F1F5F9")
            cl.font = Font(name=FN, bold=True, size=11, color="FFFFFF" if hl else hx(C["text"]))
            cl.alignment, cl.border = center, gapb
        r += 1
        si = 0
        for kind, lab, s, e in PLAN:
            if kind == "slot":
                ws.row_dimensions[r].height = 46
                tl = ws.cell(r, 1, f"{ordinal(si)}\n{s}-{e}")
                tl.font, tl.alignment = Font(name=FN, bold=True, size=9, color=hx(C["muted"])), center
                for d in range(len(DAYS)):
                    cl = ws.cell(r, d + 2)
                    cl.border, cl.alignment = gapb, left
                    xc = t["cells"].get((d, si))
                    if xc:
                        bg, fg = _colors(xc["key"])
                        l1, l2 = _lines(t["mode"], xc)
                        cl.fill = fill(bg)
                        if CellRichText is not None:
                            parts = [TextBlock(InlineFont(rFont=FN, sz=11, b=True, color=hx(fg)), l1)]
                            if l2:
                                parts.append(TextBlock(InlineFont(rFont=FN, sz=9, color=hx(fg)), "\n" + l2))
                            cl.value = CellRichText(*parts)
                        else:
                            cl.value = l1 + (f"\n{l2}" if l2 else "")
                            cl.font = Font(name=FN, bold=True, size=10, color=hx(fg))
                    else:
                        cl.fill = fill(C["soft"])
                si += 1
            else:
                ws.row_dimensions[r].height = 20
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
                bar = ws.cell(r, 1, f"{lab}  ·  {s} – {e}")
                bar.fill, bar.alignment = fill("#EEF2F7"), center
                bar.font = Font(name=FN, bold=True, size=9, color=hx(C["muted"]))
                for cc in range(2, ncols + 1):
                    ws.cell(r, cc).fill = fill("#EEF2F7")
            r += 1
        ws.freeze_panes = f"B{FIRST + 1}"
        ws.print_area = f"A1:{get_column_letter(ncols)}{r - 1}"

    # ---- vue générale
    def general(t):
        ws = sheet(t["sheet"])
        ns = len(SLOTS)
        ncols = 2 + ns
        title_rows(ws, t, ncols)
        ws.column_dimensions["A"].width = 16
        ws.column_dimensions["B"].width = 12
        for i in range(ns):
            ws.column_dimensions[get_column_letter(i + 3)].width = 12
        r = FIRST
        ws.row_dimensions[r].height = 32
        heads = ["CLASSE", "JOURS"] + [f"{ordinal(i)}\n{s}-{e}" for i, (s, e) in enumerate(SLOTS)]
        for i, h in enumerate(heads, start=1):
            cl = ws.cell(r, i, h)
            cl.fill, cl.alignment, cl.border = fill(C["primary"]), center, gapb
            cl.font = Font(name=FN, bold=True, size=10, color="FFFFFF")
        r += 1
        for name, data in t["classes"]:
            first = r
            for d, day in enumerate(DAYS):
                ws.row_dimensions[r].height = 21
                dc = ws.cell(r, 2, day)
                dc.fill, dc.alignment, dc.border = fill("#F1F5F9"), center, gapb
                dc.font = Font(name=FN, bold=True, size=9, color=hx(C["muted"]))
                for s in range(ns):
                    cl = ws.cell(r, s + 3)
                    cl.border, cl.alignment = gapb, center
                    xc = data.get((d, s))
                    if xc:
                        bg, fg = _colors(xc["key"])
                        cl.value, cl.fill = xc["code"], fill(bg)
                        cl.font = Font(name=FN, bold=True, size=9, color=hx(fg))
                    else:
                        cl.fill = fill(C["soft"])
                r += 1
            ws.merge_cells(start_row=first, start_column=1, end_row=r - 1, end_column=1)
            nc = ws.cell(first, 1, name)
            nc.fill, nc.alignment = fill("#EEF2F7"), center
            nc.font = Font(name=FN, bold=True, size=11, color=hx(C["text"]))
            for rr in range(first, r):
                ws.cell(rr, 1).border = gapb
            ws.row_dimensions[r].height = 6
            r += 1
        if t["legend"]:
            ws.cell(r, 1, "Légende").font = Font(name=FN, bold=True, size=11, color=hx(C["text"]))
            r += 1
            for code, (subj, key) in t["legend"].items():
                bg, fg = _colors(key)
                cc = ws.cell(r, 1, code)
                cc.fill, cc.alignment, cc.border = fill(bg), center, gapb
                cc.font = Font(name=FN, bold=True, size=9, color=hx(fg))
                ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=ncols)
                ws.cell(r, 2, subj).font = Font(name=FN, size=10, color=hx(C["text"]))
                ws.cell(r, 2).alignment = Alignment(vertical="center", indent=1)
                r += 1
        ws.freeze_panes = f"C{FIRST + 1}"

    for t in tables:
        (week if t["kind"] == "week" else general)(t)
    wb.save(path)
    return path