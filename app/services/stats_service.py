"""Statistiques des notes : matières, sexe, professeurs, classes, classement."""
from collections import Counter

from app.database.connection import query

SEUIL = 10  # moyenne minimale pour être « admis »
SEXES = {"M": "M", "G": "M", "H": "M", "F": "F"}

# --- en-têtes partagés entre l'UI et l'Excel
H_MAT = ["Matière", "Élèves", "Moyenne", "Moy. classe max", "Moy. classe min", "Note max", "Note min",
         "[0-6[", "[6-10[", "[10-15[", "[15-20]"]
H_MAT_SEXE = ["Matière", "Garçons Eff.", "Garçons Moy.", "Filles Eff.", "Filles Moy."]
H_SEXE = ["Sexe", "Effectif", "Admis", "Échecs", "Réussite %", "Moyenne", "Min", "Max"]
H_PROF = ["Professeur", "Matières", "Classes", "Élèves", "Moyenne", "Note min", "Note max", "Réussite %"]
H_CLASSES = ["Classe"] + [f"{p} {x}" for p in ("G", "F", "T") for x in ("Eff.", "Adm.", "%")]
H_RANK = ["Classe", "Rang", "Matricule", "Nom", "Prénoms", "Sexe", "Moyenne", "Décision"]


# ---------------------------------------------------------------- données
def note_matiere(interro, devoir, compo):
    """Note finale d'une matière sur 20.
    HYPOTHÈSE : moyenne de classe = (interro + devoir) / 2, puis (moyenne de classe + compo) / 2.
    À aligner sur grade_service.py si votre formule est différente. Les valeurs None sont ignorées."""
    cls = [v for v in (interro, devoir) if v is not None]
    mc = sum(cls) / len(cls) if cls else None
    if mc is None and compo is None:
        return None
    if mc is None:
        return float(compo)
    if compo is None:
        return float(mc)
    return (mc + compo) / 2


def fetch_grades(year, periode=None, class_id=None):
    """Dépend du schéma. Une ligne par (élève, matière, période)."""
    sql = """SELECT c.id AS cid, c.nom AS classe, m.id AS mid, m.nom AS matiere,
                    s.id AS eid, s.matricule, s.nom, s.prenom, s.sexe,
                    g.interro, g.devoir, g.compo, COALESCE(m.coefficient, 1) AS coef
             FROM grades g
             JOIN students s ON s.id = g.student_id
             JOIN classes c ON c.id = s.classe_id
             JOIN subjects m ON m.id = g.subject_id
             WHERE c.annee_id = ?"""
    params = [year]
    if periode:
        sql += " AND g.periode = ?"
        params.append(periode)
    if class_id:
        sql += " AND c.id = ?"
        params.append(class_id)
    out = []
    for r in query(sql, tuple(params)):
        note = note_matiere(r["interro"], r["devoir"], r["compo"])
        if note is None:
            continue
        out.append({"cid": r["cid"], "classe": r["classe"], "mid": r["mid"], "matiere": r["matiere"],
                    "eid": r["eid"], "matricule": r["matricule"] or "", "nom": r["nom"], "prenom": r["prenom"],
                    "sexe": r["sexe"], "interro": r["interro"], "devoir": r["devoir"], "compo": r["compo"],
                    "note": note, "coef": float(r["coef"])})
    return out


def fetch_teachers():
    """{(classe_id, matière_id): nom du professeur}. ADAPTEZ les noms de tables/colonnes à votre schéma."""
    try:
        sql = """SELECT p.nom || ' ' || COALESCE(p.prenom, '') AS prof, a.classe_id, a.subject_id
                 FROM affectations a JOIN professeurs p ON p.id = a.professeur_id"""
        return {(r["classe_id"], r["subject_id"]): r["prof"].strip() for r in query(sql)}
    except Exception:
        return {}


# ---------------------------------------------------------------- calculs
def _pct(part, total):
    return round(100 * part / total, 1) if total else 0.0


def _bande(n):
    return 0 if n < 6 else 1 if n < 10 else 2 if n < 15 else 3


def _sexe(v):
    return SEXES.get((v or "").strip().upper()[:1], "?")


def _stats(v):
    """(moyenne, min, max) arrondis."""
    if not v:
        return 0.0, 0.0, 0.0
    return round(sum(v) / len(v), 2), round(min(v), 2), round(max(v), 2)


def _add(d, s):
    for k in ["T"] + ([s["sexe"]] if s["sexe"] in ("M", "F") else []):
        d[k][0] += 1
        d[k][1] += int(s["avg"] >= SEUIL)


def _fmt(d, name):
    return {"classe": name, **{k: (v[0], v[1], _pct(v[1], v[0])) for k, v in d.items()}}


def compute(rows, teachers=None):
    teachers = teachers or {}

    # ----- par matière (+ par sexe)
    mat = {}
    for r in rows:
        m = mat.setdefault(r["mid"], {"matiere": r["matiere"], "notes": [], "cls": {}, "sx": {"M": [], "F": []}})
        m["notes"].append(r["note"])
        m["cls"].setdefault(r["classe"], []).append(r["note"])
        sx = _sexe(r["sexe"])
        if sx in m["sx"]:
            m["sx"][sx].append(r["note"])
    matieres, matieres_sexe = [], []
    for m in sorted(mat.values(), key=lambda x: x["matiere"]):
        moy, mn, mx = _stats(m["notes"])
        cls_avg = [sum(v) / len(v) for v in m["cls"].values()]
        bandes = [0] * 4
        for v in m["notes"]:
            bandes[_bande(v)] += 1
        matieres.append({"matiere": m["matiere"], "effectif": len(m["notes"]), "moyenne": moy,
                         "moy_max": round(max(cls_avg), 2), "moy_min": round(min(cls_avg), 2),
                         "max": mx, "min": mn, "bandes": bandes})
        g, f = m["sx"]["M"], m["sx"]["F"]
        matieres_sexe.append({"matiere": m["matiere"], "M": (len(g), _stats(g)[0]), "F": (len(f), _stats(f)[0])})

    # ----- moyenne générale de chaque élève (pondérée par les coefficients)
    students = {}
    for r in rows:
        s = students.setdefault(r["eid"], {"classe": r["classe"], "matricule": r["matricule"], "nom": r["nom"],
                                           "prenom": r["prenom"], "sexe": _sexe(r["sexe"]), "pts": 0.0,
                                           "coef": 0.0})
        s["pts"] += r["note"] * r["coef"]
        s["coef"] += r["coef"]
    for s in students.values():
        s["avg"] = round(s["pts"] / s["coef"], 2) if s["coef"] else 0.0

    # ----- résultats par sexe
    sexes = []
    for k, lab in (("M", "Garçons"), ("F", "Filles"), ("T", "Ensemble")):
        lst = [s["avg"] for s in students.values() if k == "T" or s["sexe"] == k]
        if not lst:
            continue
        adm = sum(a >= SEUIL for a in lst)
        moy, mn, mx = _stats(lst)
        sexes.append({"label": lab, "effectif": len(lst), "admis": adm, "echecs": len(lst) - adm,
                      "taux": _pct(adm, len(lst)), "moyenne": moy, "min": mn, "max": mx})

    # ----- par professeur
    profs = {}
    for r in rows:
        p = teachers.get((r["cid"], r["mid"]))
        if not p:
            continue
        d = profs.setdefault(p, {"notes": [], "mat": set(), "cls": set(), "eleves": set()})
        d["notes"].append(r["note"])
        d["mat"].add(r["matiere"])
        d["cls"].add(r["classe"])
        d["eleves"].add(r["eid"])
    professeurs = []
    for p in sorted(profs):
        d = profs[p]
        moy, mn, mx = _stats(d["notes"])
        professeurs.append({"prof": p, "matieres": ", ".join(sorted(d["mat"])),
                            "classes": ", ".join(sorted(d["cls"])), "eleves": len(d["eleves"]),
                            "moyenne": moy, "min": mn, "max": mx,
                            "taux": _pct(sum(n >= SEUIL for n in d["notes"]), len(d["notes"]))})

    # ----- réussite par classe et par sexe
    cls, tot = {}, {"M": [0, 0], "F": [0, 0], "T": [0, 0]}
    for s in students.values():
        _add(cls.setdefault(s["classe"], {"M": [0, 0], "F": [0, 0], "T": [0, 0]}), s)
        _add(tot, s)
    classes = [_fmt(cls[k], k) for k in sorted(cls)]
    total = _fmt(tot, "Ensemble")

    # ----- classement par classe (ex aequo : même rang)
    ranking, by_class = [], {}
    for s in students.values():
        by_class.setdefault(s["classe"], []).append(s)
    for classe in sorted(by_class):
        lst = sorted(by_class[classe], key=lambda s: (-s["avg"], s["nom"], s["prenom"]))
        ties = Counter(s["avg"] for s in lst)
        rank, prev = 0, None
        for i, s in enumerate(lst, 1):
            if s["avg"] != prev:
                rank, prev = i, s["avg"]
            ranking.append({"classe": classe, "rang": f"{rank} ex" if ties[s["avg"]] > 1 else str(rank),
                            "matricule": s["matricule"], "nom": s["nom"], "prenom": s["prenom"],
                            "sexe": s["sexe"] if s["sexe"] != "?" else "", "moyenne": s["avg"],
                            "decision": "Admis" if s["avg"] >= SEUIL else "Échec"})

    # ----- synthèse
    n = len(students)
    admis = sum(1 for s in students.values() if s["avg"] >= SEUIL)
    avgs = [s["avg"] for s in students.values()]

    def taux(lab):
        return next((x["taux"] for x in sexes if x["label"] == lab), 0.0)

    summary = {"eleves": n, "admis": admis, "echecs": n - admis, "reussite": _pct(admis, n),
               "moyenne": round(sum(avgs) / n, 2) if n else 0.0,
               "max": max(avgs) if avgs else 0.0, "min": min(avgs) if avgs else 0.0,
               "taux_M": taux("Garçons"), "taux_F": taux("Filles")}
    return {"summary": summary, "matieres": matieres, "matieres_sexe": matieres_sexe, "sexes": sexes,
            "professeurs": professeurs, "classes": classes, "total": total, "ranking": ranking}


# ---------------------------------------------------------------- lignes de tableaux (UI + Excel)
def rows_matieres(d):
    return [(x["matiere"], x["effectif"], x["moyenne"], x["moy_max"], x["moy_min"], x["max"], x["min"],
             *x["bandes"]) for x in d["matieres"]]


def rows_matieres_sexe(d):
    return [(x["matiere"], *x["M"], *x["F"]) for x in d["matieres_sexe"]]


def rows_sexes(d):
    return [(x["label"], x["effectif"], x["admis"], x["echecs"], x["taux"], x["moyenne"], x["min"], x["max"])
            for x in d["sexes"]]


def rows_profs(d):
    return [(x["prof"], x["matieres"], x["classes"], x["eleves"], x["moyenne"], x["min"], x["max"], x["taux"])
            for x in d["professeurs"]]


def rows_classes(d):
    rows = [(x["classe"], *x["M"], *x["F"], *x["T"]) for x in d["classes"]]
    if d["classes"]:
        t = d["total"]
        rows.append((t["classe"], *t["M"], *t["F"], *t["T"]))
    return rows


def rows_ranking(d):
    return [(x["classe"], x["rang"], x["matricule"], x["nom"], x["prenom"], x["sexe"], x["moyenne"],
             x["decision"]) for x in d["ranking"]]


# ---------------------------------------------------------------- export Excel
def export_excel(data, path, titre):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    fill = PatternFill("solid", fgColor="3C50E0")
    white = Font(bold=True, color="FFFFFF")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="D0D7E2")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def sheet(name, header, rows, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = name
        ws.append(header)
        for r in rows:
            ws.append(list(r))
        for row in ws.iter_rows():
            for c in row:
                c.border = border
                if c.row == 1:
                    c.font, c.fill, c.alignment = white, fill, center
        for idx, col in enumerate(ws.iter_cols(), 1):
            w = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[get_column_letter(idx)].width = min(max(w + 3, 9), 45)

    s = data["summary"]
    sheet("Synthèse", ["Indicateur", "Valeur"], [
        ["Filtre", titre], ["Élèves évalués", s["eleves"]], ["Admis (moy. ≥ 10)", s["admis"]],
        ["Échecs", s["echecs"]], ["Taux de réussite (%)", s["reussite"]], ["Moyenne générale", s["moyenne"]],
        ["Meilleure moyenne", s["max"]], ["Plus faible moyenne", s["min"]],
        ["Réussite garçons (%)", s["taux_M"]], ["Réussite filles (%)", s["taux_F"]]], first=True)
    sheet("Par matière", H_MAT, rows_matieres(data))
    sheet("Matière x sexe", H_MAT_SEXE, rows_matieres_sexe(data))
    sheet("Par sexe", H_SEXE, rows_sexes(data))
    sheet("Par professeur", H_PROF, rows_profs(data))
    sheet("Classe x sexe", H_CLASSES, rows_classes(data))
    sheet("Classement", H_RANK, rows_ranking(data))
    wb.save(path)