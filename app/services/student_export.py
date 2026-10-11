"""Export des listes de classe (élèves par classe) en PDF (reportlab) et Excel (openpyxl).

Même identité visuelle que l'export des emplois du temps : couleurs et police de l'application, fond de page
gris-bleu, bandeau sombre avec filet primaire, carte blanche arrondie, en-tête de tableau en couleur primaire,
lignes séparées par un filet léger, pastilles pour le sexe et le statut « Redoublant ».

Utilisation : tables, base = build_tables(scope, cid, cls_map) puis write_file(path, tables, "pdf" | "xlsx")
  scope = "classe" (une classe, identifiant cid) ou "toutes" (une feuille / page par classe).
"""
import datetime
import os
import re

from app.constants import C, FONT, PALETTE
from app.database.connection import query
from app.database.queries import cur_year
from app.repositories import student_repository
from app.services.settings_service import get_settings

HEAD = ["N°", "Matricule", "Nom", "Prénoms", "Sexe", "Statut", "Date de naissance", "Tuteur / Parent", "Téléphone"]
WEIGHTS = [26, 70, 100, 130, 40, 74, 84, 120, 80]
G200 = "#E4E7EC"
WARN = ("#FFFAEB", "#DC6803")                      # badge « Redoublant » (comme dans les tableaux)
SEX_STYLE = {"M": PALETTE[0], "F": PALETTE[3]}


# ----------------------------------------------------------------------------- données
def _safe(name):
    return re.sub(r"[^\w\-]+", "_", name).strip("_") or "classe"


def _g(row, i):
    """Valeur i de la ligne (tuple de LIST_SQL : id, matricule, nom, prénoms, sexe, statut, naissance, classe, tuteur, tel)."""
    try:
        v = row[i]
    except (IndexError, KeyError):
        return ""
    return "" if v is None else str(v)


def _students(year, class_id):
    rows = [tuple(r) for r in query(student_repository.LIST_SQL, (year, class_id, class_id))]
    out = [[_g(r, 1), _g(r, 2), _g(r, 3), _g(r, 4), _g(r, 5), _g(r, 6), _g(r, 8), _g(r, 9)] for r in rows]
    out.sort(key=lambda s: (s[1].lower(), s[2].lower()))
    return out


def _sex(s):
    return "F" if s.strip().lower().startswith("f") else "M" if s.strip().lower().startswith("m") else ""


def _table(name, students):
    n = len(students)
    m = sum(1 for s in students if _sex(s[3]) == "M")
    f = sum(1 for s in students if _sex(s[3]) == "F")
    r = sum(1 for s in students if s[4].lower().startswith("red"))
    sub = f"Effectif : {n} élève{'s' if n > 1 else ''}  ·  {m} garçon{'s' if m > 1 else ''}  ·  {f} fille{'s' if f > 1 else ''}"
    if r:
        sub += f"  ·  {r} redoublant{'s' if r > 1 else ''}"
    return dict(title=f"Liste de classe — {name}", sheet=name, subtitle=sub, students=students)


def build_tables(scope, cid, cls_map):
    """-> (liste de tableaux, nom de fichier de base). Liste vide s'il n'y a rien à exporter."""
    year = cur_year()
    items = list(cls_map.items()) if scope == "toutes" else [(n, i) for n, i in cls_map.items() if i == cid]
    tables = []
    for name, i in items:
        students = _students(year, i)
        if students or scope != "toutes":
            tables.append(_table(name, students))
    base = "Listes_de_classe" if scope == "toutes" else "Liste_" + _safe(items[0][0] if items else "classe")
    return tables, base


def write_file(path, tables, fmt):
    if fmt == "pdf":
        return _write_pdf(path, tables)
    if fmt == "xlsx":
        return _write_xlsx(path, tables)
    raise ValueError(f"Format non pris en charge : {fmt}")


# ============================================================================= PDF
def _fonts():
    """Segoe UI (police de l'application) si présente, sinon Helvetica."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    d = os.environ.get("WINDIR", r"C:\Windows") + r"\Fonts"
    reg, bold = os.path.join(d, "segoeui.ttf"), os.path.join(d, "segoeuib.ttf")
    try:
        if os.path.exists(reg) and os.path.exists(bold):
            pdfmetrics.registerFont(TTFont("SegoeUI", reg))
            pdfmetrics.registerFont(TTFont("SegoeUI-Bold", bold))
            return "SegoeUI", "SegoeUI-Bold"
    except Exception:
        pass
    return "Helvetica", "Helvetica-Bold"


def _write_pdf(path, tables):
    from reportlab.lib.colors import HexColor
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.pdfgen import canvas

    F, FB = _fonts()
    PW, PH = landscape(A4)
    M, PAD, BANNER_H = 28, 16, 40
    X0 = M + PAD
    W = PW - 2 * X0
    S = get_settings()
    today = datetime.date.today().strftime("%d/%m/%Y")
    col = lambda h: HexColor(h)
    TEXT, MUTED, BORDER = col(C["text"]), col(C["muted"]), col(G200)
    PRIMARY, SIDEBAR = col(C["primary"]), col(C["sidebar"])

    card_top, card_bottom = PH - BANNER_H - 14, M
    top_y = card_top - 58
    ROW_H, HEAD_H = 22, 26
    per_page = max(1, int((top_y - (card_bottom + 14) - HEAD_H - 4) // ROW_H))
    k = W / sum(WEIGHTS)
    widths = [w * k for w in WEIGHTS]
    xs = [X0 + sum(widths[:i]) for i in range(len(widths))]

    def fit(text, font, size, maxw):
        text = str(text or "")
        if stringWidth(text, font, size) <= maxw:
            return text
        while text and stringWidth(text + "…", font, size) > maxw:
            text = text[:-1]
        return text + "…"

    def rect(x, y, w, h, bg, edge, r=7):
        c.setFillColor(col(bg) if isinstance(bg, str) else bg)
        c.setStrokeColor(col(edge) if isinstance(edge, str) else edge)
        c.setLineWidth(0.7)
        c.roundRect(x, y, w, h, r, fill=1, stroke=1)

    pages = []
    for t in tables:
        rows = t["students"]
        chunks = [rows[i:i + per_page] for i in range(0, len(rows), per_page)] or [[]]
        for ci, ch in enumerate(chunks):
            pages.append((t, ci * per_page, ch))
    total = len(pages)

    c = canvas.Canvas(path, pagesize=(PW, PH), pageCompression=1)
    c.setTitle(tables[0]["title"] if len(tables) == 1 else "Listes de classe")
    c.setAuthor("EduManager")

    for no, (t, start, chunk) in enumerate(pages, start=1):
        # fond, bandeau, filet, carte
        c.setFillColor(col(C["bg"]))
        c.rect(0, 0, PW, PH, fill=1, stroke=0)
        c.setFillColor(SIDEBAR)
        c.rect(0, PH - BANNER_H, PW, BANNER_H, fill=1, stroke=0)
        c.setFillColor(PRIMARY)
        c.rect(0, PH - BANNER_H - 2.5, PW, 2.5, fill=1, stroke=0)
        c.setFillColor(col("#FFFFFF"))
        c.setFont(FB, 11)
        c.drawString(M, PH - BANNER_H / 2 - 4, str(S["ecole_nom"]).upper())
        c.setFillColor(col(C["sidebar_text"]))
        c.setFont(F, 9)
        c.drawRightString(PW - M, PH - BANNER_H / 2 - 3, f"Année scolaire {S['annee']}")
        rect(M, card_bottom, PW - 2 * M, card_top - card_bottom, C["card"], C["border"], 12)
        # titre
        c.setFillColor(TEXT)
        c.setFont(FB, 17)
        c.drawString(X0, card_top - 28, fit(t["title"], FB, 17, W))
        c.setFillColor(MUTED)
        c.setFont(F, 10)
        c.drawString(X0, card_top - 43, t["subtitle"])
        # en-tête du tableau
        y = top_y
        for i, h in enumerate(HEAD):
            rect(xs[i] + 1.5, y - HEAD_H, widths[i] - 3, HEAD_H - 2, PRIMARY, PRIMARY, 5)
            c.setFillColor(col("#FFFFFF"))
            c.setFont(FB, 8.5)
            if i in (0, 4):
                c.drawCentredString(xs[i] + widths[i] / 2, y - HEAD_H / 2 - 3, fit(h, FB, 8.5, widths[i] - 6))
            else:
                c.drawString(xs[i] + 8, y - HEAD_H / 2 - 3, fit(h, FB, 8.5, widths[i] - 12))
        y -= HEAD_H + 4
        # lignes
        if not chunk:
            c.setFillColor(MUTED)
            c.setFont(F, 11)
            c.drawCentredString(PW / 2, y - 40, "Aucun élève dans cette classe")
        for r, s in enumerate(chunk):
            ytop = y - r * ROW_H
            base = ytop - ROW_H / 2 - 3
            c.setFillColor(MUTED)
            c.setFont(F, 8.5)
            c.drawCentredString(xs[0] + widths[0] / 2, base, str(start + r + 1))
            cells = {1: s[0], 2: s[1], 3: s[2], 5: s[5], 6: s[6], 7: s[7]}
            for i, txt in cells.items():
                bold = i == 2
                c.setFillColor(TEXT if i in (2, 3) else col(C["muted"]) if i == 1 else TEXT)
                c.setFont(FB if bold else F, 9)
                c.drawString(xs[i] + 8, base, fit(txt, FB if bold else F, 9, widths[i] - 12))
            sx = _sex(s[3])                                           # pastille sexe
            if sx:
                bg, fg = SEX_STYLE[sx]
                rect(xs[4] + widths[4] / 2 - 9, ytop - ROW_H + 4, 18, ROW_H - 8, bg, bg, 6)
                c.setFillColor(col(fg))
                c.setFont(FB, 8)
                c.drawCentredString(xs[4] + widths[4] / 2, base, sx)
            if s[4].lower().startswith("red"):                        # pastille redoublant
                bg, fg = WARN
                rect(xs[5] + 6, ytop - ROW_H + 4, min(60, widths[5] - 10), ROW_H - 8, bg, bg, 7)
                c.setFillColor(col(fg))
                c.setFont(FB, 8)
                c.drawString(xs[5] + 12, base, "Redoublant")
            else:
                c.setFillColor(MUTED)
                c.setFont(F, 9)
                c.drawString(xs[5] + 8, base, fit(s[4], F, 9, widths[5] - 12))
            c.setStrokeColor(BORDER)
            c.setLineWidth(0.6)
            c.line(X0, ytop - ROW_H, X0 + W, ytop - ROW_H)
        # pied de page
        c.setFillColor(MUTED)
        c.setFont(F, 7.5)
        c.drawString(M, M - 14, f"Édité le {today}")
        c.drawRightString(PW - M, M - 14, f"Page {no} / {total}")
        c.showPage()
    c.save()
    return path


# ============================================================================= Excel
def _write_xlsx(path, tables):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.properties import PageSetupProperties

    S = get_settings()
    hx = lambda h: h.lstrip("#").upper()
    fill = lambda h: PatternFill("solid", fgColor=hx(h))
    rowb = Border(bottom=Side(style="thin", color=hx(G200)))
    gapb = Border(left=Side(style="medium", color="FFFFFF"), right=Side(style="medium", color="FFFFFF"))
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    ncols = len(HEAD)
    wb = Workbook()
    wb.remove(wb.active)
    used = set()

    for t in tables:
        base = re.sub(r"[\[\]\*\?/\\:]", " ", t["sheet"] or "Classe").strip()[:28] or "Classe"
        name, n = base, 2
        while name.lower() in used:
            name, n = f"{base[:26]} {n}", n + 1
        used.add(name.lower())
        ws = wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.page_margins.top = ws.page_margins.bottom = 0.5
        ws.sheet_properties.tabColor = hx(C["primary"])
        for i, w in enumerate(WEIGHTS, start=1):
            ws.column_dimensions[get_column_letter(i)].width = max(6, round(w / 5))

        # bandeau, filet, titre, sous-titre
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
        b = ws.cell(1, 1, f"{str(S['ecole_nom']).upper()}   ·   Année scolaire {S['annee']}")
        b.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        b.alignment = Alignment(vertical="center", indent=1)
        for cc in range(1, ncols + 1):
            ws.cell(1, cc).fill = fill(C["sidebar"])
            ws.cell(2, cc).fill = fill(C["primary"])
        ws.row_dimensions[1].height, ws.row_dimensions[2].height = 26, 3
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=ncols)
        ws.cell(3, 1, t["title"]).font = Font(name=FONT, size=16, bold=True, color=hx(C["text"]))
        ws.cell(3, 1).alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[3].height = 32
        ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=ncols)
        ws.cell(4, 1, t["subtitle"]).font = Font(name=FONT, size=10, color=hx(C["muted"]))
        ws.cell(4, 1).alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[4].height, ws.row_dimensions[5].height = 16, 8

        # en-tête du tableau
        r = 6
        ws.row_dimensions[r].height = 26
        for i, h in enumerate(HEAD, start=1):
            cl = ws.cell(r, i, h)
            cl.fill, cl.border = fill(C["primary"]), gapb
            cl.font = Font(name=FONT, bold=True, size=10, color="FFFFFF")
            cl.alignment = center if i in (1, 5) else left
        r += 1
        # lignes
        for idx, s in enumerate(t["students"], start=1):
            ws.row_dimensions[r].height = 21
            vals = [idx, s[0], s[1], s[2], _sex(s[3]) or s[3], s[4], s[5], s[6], s[7]]
            for i, v in enumerate(vals, start=1):
                cl = ws.cell(r, i, v)
                cl.border = rowb
                cl.alignment = center if i in (1, 5) else left
                cl.font = Font(name=FONT, size=10, bold=(i == 3), color=hx(C["muted"] if i in (1, 2) else C["text"]))
            sx = _sex(s[3])
            if sx:
                bg, fg = SEX_STYLE[sx]
                ws.cell(r, 5).fill = fill(bg)
                ws.cell(r, 5).font = Font(name=FONT, size=10, bold=True, color=hx(fg))
            if s[4].lower().startswith("red"):
                ws.cell(r, 6).fill = fill(WARN[0])
                ws.cell(r, 6).font = Font(name=FONT, size=10, bold=True, color=hx(WARN[1]))
            r += 1
        if not t["students"]:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
            ws.cell(r, 1, "Aucun élève dans cette classe").font = Font(name=FONT, size=11, color=hx(C["muted"]))
            ws.cell(r, 1).alignment = center
            r += 1
        ws.freeze_panes = "A7"
        ws.print_title_rows = "6:6"
        ws.print_area = f"A1:{get_column_letter(ncols)}{r - 1}"

    wb.save(path)
    return path