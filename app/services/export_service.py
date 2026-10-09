
"""Export des emplois du temps en PDF natif et Excel."""
from pathlib import Path

from app.constants import DAYS, PLAN
from app.services.schedule_service import cell_data
from app.services.settings_service import get_settings



def build_tables(scope, cid, tid, cls_map, teachers):
    """Construit les calendriers correspondant à la portée sélectionnée."""
    settings = get_settings()
    school = settings["ecole_nom"]
    year = settings["annee"]
    tables = []

    if scope == "classe":
        ids = [cid] if cid else []
        for class_id in ids:
            table = _calendar_table(
                "Classe", class_id, cls_map, teachers, school, year
            )
            if table:
                tables.append(table)

    elif scope == "enseignant":
        ids = [tid] if tid else []
        for teacher_id in ids:
            table = _calendar_table(
                "Enseignant", teacher_id, cls_map, teachers, school, year
            )
            if table:
                tables.append(table)

    elif scope in ("general", "classes"):
        for class_id in cls_map.values():
            table = _calendar_table(
                "Classe", class_id, cls_map, teachers, school, year
            )
            if table:
                tables.append(table)

    elif scope == "enseignants":
        for teacher_id in teachers.values():
            table = _calendar_table(
                "Enseignant", teacher_id, cls_map, teachers, school, year
            )
            if table:
                tables.append(table)

    else:
        raise ValueError(f"Portée inconnue : {scope}")

    if not tables:
        return [], "Emploi_du_temps"

    if len(tables) == 1:
        base = tables[0]["sheet"]
    else:
        base = {
            "general": "Emploi_du_temps_general",
            "classes": "Emplois_du_temps_classes",
            "enseignants": "Emplois_du_temps_enseignants",
        }.get(scope, "Emplois_du_temps")

    return tables, base

def _calendar_table(mode, entity_id, cls_map, teachers, school, year):
    """Produit une grille identique à la structure du calendrier."""
    if mode == "Classe":
        name = next(
            (n for n, i in cls_map.items() if i == entity_id), None
        )
        if name is None:
            return None
        data = cell_data("c", entity_id)
        title = f"Emploi du temps — {name}"
    else:
        name = next(
            (n for n, i in teachers.items() if i == entity_id), None
        )
        if name is None:
            return None
        data = cell_data("t", entity_id)
        title = f"Emploi du temps — {name}"

    rows = []
    merges = []
    slot_index = 0

    # Chaque ligne de PLAN correspond à un créneau ou à une pause.
    for kind, label, start, end in PLAN:
        if kind == "slot":
            cells = []
            for day_index in range(len(DAYS)):
                item = data.get((day_index, slot_index))
                if not item:
                    cells.append("")
                elif mode == "Classe":
                    subject = item.get("subj", "")
                    teacher = item.get("teacher", "")
                    cells.append(
                        f"{subject}\n{teacher}" if teacher else subject
                    )
                else:
                    class_name = item.get("cls", "")
                    subject = item.get("subj", "")
                    cells.append(f"{class_name}\n{subject}")

            rows.append({
                "type": "slot",
                "time": f"{start} - {end}",
                "label": f"{start} - {end}",
                "cells": cells,
            })
            slot_index += 1
        else:
            rows.append({
                "type": "break",
                "time": f"{start} - {end}",
                "label": f"{label}  ·  {start} - {end}",
                "cells": [label] * len(DAYS),
            })
            merges.append(len(rows))

    return {
        "title": title,
        "sheet": name[:31],
        "caption": f"{school} — Année scolaire {year}",
        "header": ["HORAIRES"] + list(DAYS),
        "rows": rows,
        "merges": merges,
    }


def write_file(path, tables, fmt_):
    """Écrit un vrai PDF ou un classeur Excel."""
    if fmt_ == "xlsx":
        _write_excel(path, tables)
    elif fmt_ == "pdf":
        _write_pdf(path, tables)
    else:
        raise ValueError(f"Format d'export non pris en charge : {fmt_}")


def _write_excel(path, tables):
    from openpyxl import Workbook
    from openpyxl.styles import (
        Alignment, Border, Font, PatternFill, Side
    )
    from openpyxl.worksheet.page import PageMargins

    wb = Workbook()
    wb.remove(wb.active)

    navy = "233B67"
    blue = "3C50E0"
    light_blue = "EEF2FF"
    light_gray = "F1F5F9"
    white = "FFFFFF"
    border_side = Side(style="thin", color="CBD5E1")
    border = Border(
        left=border_side, right=border_side,
        top=border_side, bottom=border_side,
    )

    for index, table in enumerate(tables, start=1):
        ws = wb.create_sheet(
            title=(table["sheet"] or f"Calendrier {index}")[:31]
        )
        ws.sheet_view.showGridLines = False

        last_col = len(DAYS) + 1
        ws.merge_cells(
            start_row=1, start_column=1,
            end_row=1, end_column=last_col,
        )
        title_cell = ws.cell(1, 1, table["title"])
        title_cell.font = Font(
            name="Arial", size=16, bold=True, color=white
        )
        title_cell.fill = PatternFill("solid", fgColor=navy)
        title_cell.alignment = Alignment(
            horizontal="center", vertical="center"
        )
        ws.row_dimensions[1].height = 32

        ws.merge_cells(
            start_row=2, start_column=1,
            end_row=2, end_column=last_col,
        )
        ws.cell(2, 1, table["caption"]).font = Font(
            name="Arial", size=10, italic=True, color="64748B"
        )
        ws.cell(2, 1).alignment = Alignment(horizontal="center")
        ws.row_dimensions[2].height = 23

        header_row = 4
        for col, value in enumerate(table["header"], start=1):
            cell = ws.cell(header_row, col, value)
            cell.font = Font(bold=True, color=white, size=10)
            cell.fill = PatternFill("solid", fgColor=blue)
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            cell.border = border
        ws.row_dimensions[header_row].height = 28

        for row_index, item in enumerate(table["rows"], start=5):
            ws.cell(row_index, 1, item["label"])
            for col, value in enumerate(item["cells"], start=2):
                ws.cell(row_index, col, value)

            is_break = item["type"] == "break"
            if is_break:
                ws.merge_cells(
                    start_row=row_index, start_column=2,
                    end_row=row_index, end_column=last_col,
                )
                ws.cell(row_index, 2, item["label"])
            for col in range(1, last_col + 1):
                cell = ws.cell(row_index, col)
                cell.border = border
                cell.alignment = Alignment(
                    horizontal="center", vertical="center",
                    wrap_text=True,
                )
                if is_break:
                    cell.fill = PatternFill("solid", fgColor=light_gray)
                    cell.font = Font(
                        bold=True, color="475569", size=9
                    )
                else:
                    cell.fill = PatternFill(
                        "solid",
                        fgColor=light_blue if row_index % 2 else white,
                    )
                    cell.font = Font(size=10, color="1E293B")
            ws.row_dimensions[row_index].height = 43 if not is_break else 23

        ws.column_dimensions["A"].width = 17
        for col in range(2, last_col + 1):
            ws.column_dimensions[
                ws.cell(1, col).column_letter
            ].width = 22

        ws.freeze_panes = "B5"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.page_margins = PageMargins(
            left=0.25, right=0.25, top=0.45, bottom=0.45,
            header=0.2, footer=0.2,
        )
        ws.print_title_rows = "1:4"
        ws.print_area = f"A1:{ws.cell(4 + len(table['rows']), last_col).coordinate}"

    wb.save(path)


def _write_pdf(path, tables):
    """Produit un PDF natif avec ReportLab, sans conversion HTML."""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
    )
    from xml.sax.saxutils import escape

    navy = colors.HexColor("#233B67")
    blue = colors.HexColor("#3C50E0")
    pale = colors.HexColor("#EEF2FF")
    gray = colors.HexColor("#F1F5F9")
    border = colors.HexColor("#CBD5E1")

    doc = SimpleDocTemplate(
        str(path),
        pagesize=landscape(A4),
        rightMargin=8 * mm, leftMargin=8 * mm,
        topMargin=10 * mm, bottomMargin=10 * mm,
        title="Emplois du temps",
        author="EduManager",
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "CalendarTitle", parent=styles["Title"],
        fontName="Helvetica-Bold", fontSize=16, leading=20,
        textColor=navy, alignment=TA_CENTER, spaceAfter=4,
    )
    caption_style = ParagraphStyle(
        "CalendarCaption", parent=styles["Normal"],
        fontSize=9, leading=12, textColor=colors.HexColor("#64748B"),
        alignment=TA_CENTER, spaceAfter=10,
    )
    cell_style = ParagraphStyle(
        "CalendarCell", parent=styles["Normal"],
        fontName="Helvetica", fontSize=8, leading=10,
        alignment=TA_CENTER, textColor=colors.HexColor("#1E293B"),
    )
    header_style = ParagraphStyle(
        "CalendarHeader", parent=cell_style,
        fontName="Helvetica-Bold", textColor=colors.white,
    )
    break_style = ParagraphStyle(
        "CalendarBreak", parent=cell_style,
        fontName="Helvetica-Bold", textColor=colors.HexColor("#475569"),
    )

    story = []
    usable_width = landscape(A4)[0] - 16 * mm
    time_width = 24 * mm
    day_width = (usable_width - time_width) / len(DAYS)

    for ti, item in enumerate(tables):
        story.append(Paragraph(escape(item["title"]), title_style))
        story.append(Paragraph(escape(item["caption"]), caption_style))

        matrix = [[
            Paragraph("HORAIRES", header_style),
            *[Paragraph(escape(day), header_style) for day in DAYS],
        ]]
        break_rows = []

        for index, row in enumerate(item["rows"], start=1):
            if row["type"] == "break":
                matrix.append([
                    Paragraph(escape(row["time"]), break_style),
                    Paragraph(escape(row["label"]), break_style),
                ] + [""] * (len(DAYS) - 1))
                break_rows.append(index)
            else:
                matrix.append([
                    Paragraph(escape(row["time"]), cell_style),
                    *[
                        Paragraph(
                            escape(text).replace("\n", "<br/>"),
                            cell_style,
                        )
                        for text in row["cells"]
                    ],
                ])

        grid = Table(
            matrix,
            colWidths=[time_width] + [day_width] * len(DAYS),
            repeatRows=1,
            hAlign="CENTER",
        )
        commands = [
            ("BACKGROUND", (0, 0), (-1, 0), blue),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.5, border),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]
        for r in range(1, len(matrix)):
            if r not in break_rows:
                commands.append((
                    "BACKGROUND", (0, r), (-1, r),
                    pale if r % 2 else colors.white,
                ))
        for r in break_rows:
            commands.extend([
                ("SPAN", (1, r), (-1, r)),
                ("BACKGROUND", (0, r), (-1, r), gray),
            ])
        grid.setStyle(TableStyle(commands))
        story.append(grid)

        if ti < len(tables) - 1:
            story.append(PageBreak())

    doc.build(story)