"""Import des notes depuis un fichier Excel + génération d'un modèle à remplir.

Format accepté (la première ligne contenant des en-têtes reconnus est utilisée, parmi les 20 premières) :
  - élève : une colonne « Nom » + « Prénom », ou une colonne unique « Élève » / « Nom et prénoms » ;
  - notes : « Interrogation » (ou Interro), « Devoir », « Composition » (ou Compo).
Les colonnes absentes sont ignorées ; une cellule vide, « ABS » ou « - » ne modifie pas la note existante.
"""
import re
import unicodedata

from app.utils.validation import parse_note

FIELDS = ("interro", "devoir", "compo")
LABELS = {"interro": "Interrogation", "devoir": "Devoir", "compo": "Composition"}
_SKIP = {"", "abs", "absent", "absente", "nc", "na"}
_FULL = {"eleve", "eleves", "nomprenom", "nomprenoms", "nometprenom", "nometprenoms", "nomcomplet",
         "nomsetprenoms", "nomsprenoms"}


def _norm(s):
    s = unicodedata.normalize("NFD", str(s or "")).lower()
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def _tokens(s):
    return tuple(sorted(_norm(s).split()))


def _classify(header):
    n = _norm(header).replace(" ", "")
    if not n:
        return None
    if n in ("nom", "noms"):
        return "nom"
    if n in ("prenom", "prenoms"):
        return "prenom"
    if n in _FULL:
        return "full"
    if "interro" in n or n in ("int", "inter"):
        return "interro"
    if "devoir" in n or n in ("dev", "devo"):
        return "devoir"
    if "compo" in n or n == "comp":
        return "compo"
    return None


def _cell(v):
    if v is None:
        return ""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return f"{v:g}"
    return str(v).strip().replace(",", ".")


def _find_header(rows):
    for i, row in enumerate(rows[:20]):
        cols = {}
        for j, v in enumerate(row):
            k = _classify(v)
            if k and k not in cols.values():
                cols[j] = k
        has_name = "full" in cols.values() or "nom" in cols.values()
        if has_name and any(k in FIELDS for k in cols.values()):
            return i, cols
    return None, None


def _match(name_tokens, exact, students):
    """Renvoie (id, None) si trouvé, (None, 'ambigu') si plusieurs élèves, (None, None) sinon."""
    ids = exact.get(name_tokens, [])
    if len(ids) == 1:
        return ids[0], None
    if len(ids) > 1:
        return None, "ambigu"
    cand = [sid for toks, sid in students if set(name_tokens) <= set(toks) or set(toks) <= set(name_tokens)]
    if len(cand) == 1:
        return cand[0], None
    return None, ("ambigu" if len(cand) > 1 else None)


def parse_file(path, students):
    """Lit le fichier et fait correspondre chaque ligne à un élève de la classe.

    students : liste de lignes ayant les clés id, nom, prenom.
    Retourne {"grades": {id: {champ: valeur}}, "unmatched": [...], "invalid": [...], "rows": n}.
    Lève ValueError si le fichier n'a pas le bon format, ImportError si openpyxl manque.
    """
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    h, cols = _find_header(rows)
    if h is None:
        raise ValueError("En-têtes introuvables. Le fichier doit contenir une colonne « Nom » (et « Prénom ») "
                         "ou « Élève », puis au moins une colonne parmi « Interrogation », « Devoir », "
                         "« Composition ».\n\nUtilisez le bouton « Modèle Excel » pour obtenir un fichier prêt à remplir.")
    stud = [(_tokens(f"{s['nom']} {s['prenom']}"), s["id"]) for s in students]
    exact = {}
    for toks, sid in stud:
        exact.setdefault(toks, []).append(sid)
    by_key = {k: j for j, k in cols.items()}

    out = dict(grades={}, unmatched=[], invalid=[], rows=0)
    for n, row in enumerate(rows[h + 1:], start=h + 2):
        get = lambda k: _cell(row[by_key[k]]) if k in by_key and by_key[k] < len(row) else ""
        name = get("full") or f"{get('nom')} {get('prenom')}".strip()
        values = {k: get(k) for k in FIELDS if k in by_key}
        if not name and not any(values.values()):
            continue                                   # ligne vide
        out["rows"] += 1
        if not name:
            out["unmatched"].append(f"ligne {n} (nom manquant)")
            continue
        sid, why = _match(_tokens(name), exact, stud)
        if sid is None:
            out["unmatched"].append(f"{name}" + (" (plusieurs élèves possibles)" if why else ""))
            continue
        got = out["grades"].setdefault(sid, {})
        for k, raw in values.items():
            if _norm(raw) in _SKIP:
                continue
            ok, v = parse_note(raw)
            if not ok:
                out["invalid"].append(f"{name} — {LABELS[k]} : « {raw} »")
            elif v is not None:
                got[k] = v
    out["grades"] = {sid: g for sid, g in out["grades"].items() if g}
    return out


def write_template(path, students, title=""):
    """Crée un classeur .xlsx prêt à remplir (liste des élèves + colonnes de notes validées de 0 à 20)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Notes"
    ws["A1"] = title or "Notes"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "Saisissez les notes sur 20 (virgule ou point). Laissez vide une note absente. Ne modifiez pas les noms."
    ws["A2"].font = Font(italic=True, color="666666")
    heads = ["N°", "Nom", "Prénom"] + [LABELS[k] for k in FIELDS]
    fill = PatternFill("solid", fgColor="3C50E0")
    for j, t in enumerate(heads, start=1):
        c = ws.cell(row=4, column=j, value=t)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(horizontal="center")
    for i, s in enumerate(students, start=1):
        ws.cell(row=4 + i, column=1, value=i)
        ws.cell(row=4 + i, column=2, value=str(s["nom"]).upper())
        ws.cell(row=4 + i, column=3, value=s["prenom"])
    last = 4 + max(len(students), 1)
    dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="20", allow_blank=True,
                        errorTitle="Note invalide", error="La note doit être comprise entre 0 et 20.")
    ws.add_data_validation(dv)
    dv.add(f"D5:F{last}")
    for col, w in zip("ABCDEF", (6, 26, 28, 16, 12, 14)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    wb.save(path)
    return path