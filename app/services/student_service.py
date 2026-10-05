"""Import / modèle Excel de la liste des élèves."""
import csv
import datetime
import os

from app.database.connection import conn
from app.database.queries import cur_year
from app.repositories import student_repository
from app.utils.formatting import new_matricule, norm

IMPORT_ALIASES = dict(
    nom={"nom"}, prenom={"prenom", "prenoms"}, sexe={"sexe", "genre"}, statut={"statut"},
    date_naissance={"datedenaissance", "datenaissance", "naissance", "datenaiss"},
    tuteur={"tuteur", "parent", "tuteurparent", "parents"}, telephone={"telephone", "tel", "contact"})


class StudentImportError(Exception):
    """Fichier inexploitable (colonnes manquantes...). Le message est destiné à l'utilisateur."""


def cell_text(v):
    if v is None:
        return ""
    if isinstance(v, datetime.date):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def read_table(path):
    if os.path.splitext(path)[1].lower() == ".csv":
        txt = ""
        for enc in ("utf-8-sig", "cp1252"):
            try:
                with open(path, newline="", encoding=enc) as f:
                    txt = f.read()
                break
            except UnicodeDecodeError:
                continue
        delim = ";" if txt.count(";") >= txt.count(",") else ","
        return list(csv.reader(txt.splitlines(), delimiter=delim))
    from openpyxl import load_workbook  # ImportError géré par l'appelant
    wb = load_workbook(path, read_only=True, data_only=True)
    rows = [list(r) for r in wb.active.iter_rows(values_only=True)]
    wb.close()
    return rows


def import_students(path, class_id, year=None):
    """Importe un fichier Excel/CSV dans la classe class_id (une éventuelle colonne « Classe » est ignorée).
    -> dict(added, skipped)."""
    table = read_table(path)
    hdr_i = next((i for i, r in enumerate(table) if any(norm(c) == "nom" for c in r)), None)
    if hdr_i is None:
        raise StudentImportError("Colonne « Nom » introuvable. Utilisez le « Modèle Excel » pour "
                                 "préparer votre fichier.")
    heads = [norm(c) for c in table[hdr_i]]
    idx = {k: next((i for i, h in enumerate(heads) if h in al), None) for k, al in IMPORT_ALIASES.items()}
    if idx["prenom"] is None:
        raise StudentImportError("Colonne « Prénoms » introuvable.")
    yid = year or cur_year()
    seen = student_repository.known_keys(yid)
    added = skipped = 0
    for row in table[hdr_i + 1:]:
        get = lambda k: cell_text(row[idx[k]]) if idx[k] is not None and idx[k] < len(row) else ""
        nom, prenom = get("nom"), get("prenom")
        if not nom or not prenom:
            skipped += 1
            continue
        dn = get("date_naissance")
        key = (nom.lower(), prenom.lower(), dn)
        if key in seen:
            skipped += 1
            continue
        seen.add(key)
        sx = get("sexe").lower()[:1]
        sexe = "Masculin" if sx in ("m", "g") else "Féminin" if sx == "f" else ""
        statut = "Redoublant" if get("statut").lower().startswith("red") else "Nouveau"
        sid = student_repository.create(dict(nom=nom, prenom=prenom, sexe=sexe, statut=statut, date_naissance=dn,
                                             classe_id=class_id, tuteur=get("tuteur"), telephone=get("telephone")),
                                        yid, commit=False)
        student_repository.set_matricule(sid, new_matricule(sid), commit=False)
        added += 1
    conn.commit()
    return dict(added=added, skipped=skipped)


def save_template(path):
    """Écrit le modèle Excel d'import (ImportError si openpyxl est absent)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Élèves"
    ws.append(["Nom", "Prénoms", "Sexe", "Statut", "Date de naissance", "Tuteur", "Téléphone"])
    ws.append(["KOFFI", "Ama Julie", "Féminin", "Nouveau", "12/05/2013", "KOFFI Kokou", "90000000"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="3C50E0")
    for col, w in zip("ABCDEFG", (18, 22, 12, 12, 18, 22, 14)):
        ws.column_dimensions[col].width = w
    wb.save(path)